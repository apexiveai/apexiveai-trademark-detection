
from __future__ import annotations

import difflib
import hmac
import io
import json
import os
import re
import smtplib
import time
import zipfile
from collections import deque
from pathlib import Path
from typing import Any
import cv2
import numpy as np
import pymupdf

from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, Response
from pydantic import BaseModel, Field
from openpyxl import load_workbook
from PIL import Image, ImageChops, ImageStat, UnidentifiedImageError
from pypdf import PdfReader
from starlette.concurrency import run_in_threadpool
from starlette.middleware.base import RequestResponseEndpoint
from starlette.requests import Request
from xml.etree import ElementTree

from app.auth import (
    auth_is_required,
    community_access_enabled,
    login_user,
    register_user,
    require_user,
    verify_email,
)

APP_NAME = "Apexive AI Trademark Conflict Detector"
VERSION = "1.1.0"

MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_PDF_PAGES = 500
MAX_WORKSHEETS = 200
MAX_CELLS = 250_000
MAX_TEXT_CHARS = 2_000_000
MAX_IMAGES_PER_FILE = 100
MAX_IMAGE_PIXELS = 40_000_000
MAX_RENDERED_PAGE_DIMENSION = 1600
MAX_RENDERED_SVG_DIMENSION = 512
LOGO_MATCH_THRESHOLD = 80

ALLOWED_EXTENSIONS = {
    ".pdf", ".xlsx", ".xlsm", ".txt",
    ".png", ".jpg", ".jpeg", ".webp", ".svg",
}

CORS_ORIGINS = [
    item.strip()
    for item in os.getenv(
        "CORS_ORIGINS",
        "http://localhost:5190,http://127.0.0.1:5190,https://apexiveai-trademark-detection.onrender.com,https://apexiveai-detector-v1-1.vercel.app",
    ).split(",")
    if item.strip()
]

app = FastAPI(
    title=APP_NAME,
    version=VERSION,
    description=(
        "Preliminary trademark text and image similarity screening API."
    ),
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ORIGINS,
    allow_credentials=False,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["Content-Type", "Accept", "Authorization", "X-Ingest-Key"],
)

_rate_limit_windows: dict[tuple[str, str], deque[float]] = {}
_RATE_LIMITS = {
    "/api/v1/auth/register": (5, 60),
    "/api/v1/auth/login": (12, 60),
    "/api/v1/auth/verify": (12, 60),
    "/api/v1/trademarks/search": (30, 60),
    "/api/v1/compare": (30, 60),
    "/api/v1/logo/crop": (30, 60),
    "/api/v1/trademarks": (20, 60),
}


@app.middleware("http")
async def apply_rate_limits(
    request: Request,
    call_next: RequestResponseEndpoint,
) -> Response:
    rule = _RATE_LIMITS.get(request.url.path)
    if rule is None:
        return await call_next(request)
    client_ip = request.client.host if request.client else "unknown"
    key = (client_ip, request.url.path)
    limit, window_seconds = rule
    now = time.monotonic()
    requests = _rate_limit_windows.setdefault(key, deque())
    while requests and now - requests[0] >= window_seconds:
        requests.popleft()
    if len(requests) >= limit:
        retry_after = max(1, int(window_seconds - (now - requests[0])))
        return JSONResponse(
            status_code=429,
            content={"detail": "Too many requests. Please try again shortly."},
            headers={"Retry-After": str(retry_after)},
        )
    requests.append(now)
    if len(_rate_limit_windows) > 10_000:
        expired_keys = [
            bucket
            for bucket, timestamps in _rate_limit_windows.items()
            if not timestamps or now - timestamps[-1] >= window_seconds
        ]
        for bucket in expired_keys:
            _rate_limit_windows.pop(bucket, None)
    return await call_next(request)


class AuthPayload(BaseModel):
    email: str = Field(max_length=254)
    password: str = Field(min_length=1, max_length=128)


class EmailVerificationPayload(BaseModel):
    token: str = Field(min_length=1, max_length=256)


def clean_text(value: str) -> str:
    value = value.casefold()
    value = re.sub(r"\s+", " ", value)
    return value.strip()



def extract_application_numbers(text: str) -> list[str]:
    """
    Extract explicitly labelled application numbers.
    Reject ordinary words such as 'mark', 'logo', and 'name'.
    """
    patterns = [
        r"\bapplication\s*(?:no\.?|number|#)\s*[:\-]?\s*"
        r"([A-Z0-9][A-Z0-9./\-]{2,30})",

        r"\bapp\.?\s*(?:no\.?|number|#)\s*[:\-]?\s*"
        r"([A-Z0-9][A-Z0-9./\-]{2,30})",

        r"\btrademark\s+application\s*"
        r"(?:no\.?|number|#)\s*[:\-]?\s*"
        r"([A-Z0-9][A-Z0-9./\-]{2,30})",
    ]

    ignored_labels = {
        "mark", "logo", "name", "image", "word",
        "device", "class", "date", "number", "unknown",
    }

    found: set[str] = set()

    for pattern in patterns:
        for match in re.finditer(pattern, text, re.IGNORECASE):
            candidate = match.group(1).strip(" .-/")
            normalized = candidate.casefold()

            # Application identifiers should normally contain digits.
            # Avoid treating headings and ordinary labels as IDs.
            if normalized in ignored_labels:
                continue

            if not any(character.isdigit() for character in candidate):
                continue

            if len(candidate) < 3:
                continue

            found.add(normalized)

    return sorted(found)



def remove_application_number_labels(text: str) -> str:
    """
    Remove labelled application-number fields before text scoring,
    so a matching application number alone does not inflate the score.
    """
    patterns = [
        r"\bapplication\s*(?:no\.?|number|#)\s*[:\-]?\s*"
        r"[A-Z0-9][A-Z0-9./\-]{2,30}",
        r"\bapp\.?\s*(?:no\.?|number|#)\s*[:\-]?\s*"
        r"[A-Z0-9][A-Z0-9./\-]{2,30}",
        r"\btrademark\s*(?:application\s*)?"
        r"(?:no\.?|number|#)\s*[:\-]?\s*"
        r"[A-Z0-9][A-Z0-9./\-]{2,30}",
    ]

    result = text

    for pattern in patterns:
        result = re.sub(pattern, " ", result, flags=re.IGNORECASE)

    return result


def calculate_text_similarity(a: str, b: str) -> float | None:
    a = clean_text(a)
    b = clean_text(b)

    if not a or not b:
        return None

    if a == b:
        return 100.0

    score = difflib.SequenceMatcher(
        None,
        a[:200_000],
        b[:200_000],
        autojunk=False,
    ).ratio()

    return round(score * 100, 2)


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue

    return data.decode("utf-8", errors="replace")


def validate_raster(data: bytes) -> bytes:
    try:
        with Image.open(io.BytesIO(data)) as image:
            if image.width * image.height > MAX_IMAGE_PIXELS:
                raise ValueError(
                    "Image exceeds the configured 40-megapixel limit."
                )
            image.verify()

    except UnidentifiedImageError as exc:
        raise ValueError("Invalid or unsupported image file.") from exc

    except OSError as exc:
        raise ValueError("Could not decode the image file.") from exc

    return data


def extract_svg_text(data: bytes) -> str:
    try:
        root = ElementTree.fromstring(data)
    except ElementTree.ParseError as exc:
        raise ValueError("Invalid SVG/XML file.") from exc

    parts: list[str] = []

    for node in root.iter():
        tag = node.tag.rsplit("}", 1)[-1].lower()

        # Do not execute or interpret SVG scripts.
        if tag in {"script", "foreignobject", "style"}:
            continue

        if tag in {"text", "tspan", "title", "desc"} and node.text:
            parts.append(node.text)

        for key in ("aria-label", "title", "alt"):
            value = node.attrib.get(key)
            if value:
                parts.append(value)

    return " ".join(parts)[:MAX_TEXT_CHARS]


def extract_pdf(
    data: bytes,
) -> tuple[str, list[bytes], list[str], list[str]]:
    reader = PdfReader(io.BytesIO(data), strict=False)

    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception:
            unlocked = 0

        if not unlocked:
            raise ValueError(
                "Password-protected PDF files are not supported."
            )

    total_pages = len(reader.pages)
    text_parts: list[str] = []
    images: list[bytes] = []
    image_locations: list[str] = []
    limitations: list[str] = []

    if total_pages > MAX_PDF_PAGES:
        limitations.append(
            f"PDF has {total_pages} pages; only the first "
            f"{MAX_PDF_PAGES} pages were processed."
        )

    try:
        render_document = pymupdf.open(stream=data, filetype="pdf")
    except Exception:
        render_document = None
        limitations.append(
            "PDF pages could not be rendered for visual logo comparison."
        )

    try:
        for page_index, page in enumerate(reader.pages[:MAX_PDF_PAGES]):
            embedded_image_index = 0
            try:
                page_text = page.extract_text() or ""
                if page_text:
                    text_parts.append(page_text)
            except Exception:
                limitations.append(
                    f"Text extraction failed on PDF page {page_index + 1}."
                )

            if len(images) < MAX_IMAGES_PER_FILE:
                try:
                    for embedded in page.images:
                        if len(images) >= MAX_IMAGES_PER_FILE:
                            break

                        candidate = embedded.data

                        try:
                            validate_raster(candidate)
                            images.append(candidate)
                            embedded_image_index += 1
                            image_locations.append(
                                f"Page {page_index + 1} · embedded image "
                                f"{embedded_image_index}"
                            )
                        except ValueError:
                            limitations.append(
                                "An unsupported embedded image was skipped "
                                f"on PDF page {page_index + 1}."
                            )
                except Exception:
                    limitations.append(
                        "Embedded images could not be extracted from "
                        f"PDF page {page_index + 1}."
                    )

            if (
                render_document is not None
                and page_index < len(render_document)
                and len(images) < MAX_IMAGES_PER_FILE
            ):
                try:
                    rendered_page = render_document[page_index]
                    scale = min(
                        2.0,
                        MAX_RENDERED_PAGE_DIMENSION
                        / max(rendered_page.rect.width, rendered_page.rect.height),
                    )
                    pixmap = rendered_page.get_pixmap(
                        matrix=pymupdf.Matrix(scale, scale),
                        alpha=False,
                    )
                    candidate = pixmap.tobytes("png")
                    validate_raster(candidate)
                    images.append(candidate)
                    image_locations.append(
                        f"Page {page_index + 1} · rendered page"
                    )
                except Exception:
                    limitations.append(
                        "PDF page "
                        f"{page_index + 1} could not be rendered for "
                        "visual logo comparison."
                    )

            if sum(map(len, text_parts)) >= MAX_TEXT_CHARS:
                limitations.append(
                    "PDF text extraction reached the configured text limit."
                )
                break
    finally:
        if render_document is not None:
            render_document.close()

    if len(images) >= MAX_IMAGES_PER_FILE:
        limitations.append(
            f"Visual image extraction reached the {MAX_IMAGES_PER_FILE}-image "
            "limit; additional page or embedded images may not have been "
            "compared."
        )

    return (
        "\n".join(text_parts)[:MAX_TEXT_CHARS],
        images,
        limitations,
        image_locations,
    )


def rasterize_svg(data: bytes) -> bytes:
    svg_document = None
    document = None
    try:
        svg_document = pymupdf.open(stream=data, filetype="svg")
        pdf_data = svg_document.convert_to_pdf()
        document = pymupdf.open(stream=pdf_data, filetype="pdf")
        page = document[0]
        scale = min(
            MAX_RENDERED_SVG_DIMENSION / page.rect.width,
            MAX_RENDERED_SVG_DIMENSION / page.rect.height,
        )
        rendered = page.get_pixmap(
            matrix=pymupdf.Matrix(scale, scale),
            alpha=False,
        ).tobytes("png")
        validate_raster(rendered)
        return rendered
    except Exception as exc:
        raise ValueError(
            "Could not render this SVG file for visual comparison."
        ) from exc
    finally:
        if document is not None:
            document.close()
        if svg_document is not None:
            svg_document.close()


def extract_xlsx(
    data: bytes,
) -> tuple[str, list[bytes], list[str], list[str]]:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ValueError("Invalid XLSX/XLSM document.")

    try:
        workbook = load_workbook(
            io.BytesIO(data),
            read_only=False,
            data_only=True,
        )
    except Exception as exc:
        raise ValueError(
            "Could not open the spreadsheet. Check that it is valid."
        ) from exc

    text_parts: list[str] = []
    images: list[bytes] = []
    image_locations: list[str] = []
    limitations: list[str] = []
    cell_count = 0

    try:
        worksheets = workbook.worksheets

        if len(worksheets) > MAX_WORKSHEETS:
            limitations.append(
                f"Workbook has {len(worksheets)} worksheets; only the "
                f"first {MAX_WORKSHEETS} were processed."
            )

        for worksheet_index, worksheet in enumerate(
            worksheets[:MAX_WORKSHEETS],
            start=1,
        ):
            text_parts.append(f"WORKSHEET: {worksheet.title}")

            for row in worksheet.iter_rows(values_only=True):
                for value in row:
                    cell_count += 1

                    if cell_count > MAX_CELLS:
                        limitations.append(
                            f"Spreadsheet extraction stopped at "
                            f"{MAX_CELLS:,} cells."
                        )
                        return (
                            "\n".join(text_parts)[:MAX_TEXT_CHARS],
                            images,
                            limitations,
                            image_locations,
                        )

                    if value is not None:
                        text_parts.append(str(value))

                if sum(map(len, text_parts)) >= MAX_TEXT_CHARS:
                    limitations.append(
                        "Spreadsheet text extraction reached its limit."
                    )
                    return (
                        "\n".join(text_parts)[:MAX_TEXT_CHARS],
                        images,
                        limitations,
                        image_locations,
                    )

            # openpyxl exposes embedded worksheet images through _images.
            # This is a library implementation detail and should be
            # regression-tested when upgrading openpyxl.
            try:
                for embedded_image_index, embedded in enumerate(
                    getattr(worksheet, "_images", []),
                    start=1,
                ):
                    if len(images) >= MAX_IMAGES_PER_FILE:
                        limitations.append(
                            f"Image extraction stopped at "
                            f"{MAX_IMAGES_PER_FILE} images."
                        )
                        break

                    try:
                        candidate = embedded._data()
                        validate_raster(candidate)
                        images.append(candidate)
                        image_locations.append(
                            f"Sheet {worksheet_index} · {worksheet.title} · "
                            "embedded image "
                            f"{embedded_image_index}"
                        )
                    except Exception:
                        continue
            except Exception:
                limitations.append(
                    f"Could not inspect all embedded images in "
                    f"worksheet {worksheet.title}."
                )

        return (
            "\n".join(text_parts)[:MAX_TEXT_CHARS],
            images,
            limitations,
            image_locations,
        )

    finally:
        workbook.close()


def analyze_file(filename: str, data: bytes) -> dict[str, Any]:
    ext = Path(filename).suffix.lower()

    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext or '(none)'}")

    result: dict[str, Any] = {
        "filename": filename,
        "text": "",
        "images": [],
        "image_locations": [],
        "limitations": [],
    }

    if ext == ".pdf":
        text, images, limitations, image_locations = extract_pdf(data)
        result["text"] = text
        result["images"] = images
        result["image_locations"] = image_locations
        result["limitations"].extend(limitations)

    elif ext in {".xlsx", ".xlsm"}:
        text, images, limitations, image_locations = extract_xlsx(data)
        result["text"] = text
        result["images"] = images
        result["image_locations"] = image_locations
        result["limitations"].extend(limitations)

    elif ext == ".txt":
        result["text"] = decode_text(data)[:MAX_TEXT_CHARS]

    elif ext == ".svg":
        result["text"] = extract_svg_text(data)
        result["images"] = [rasterize_svg(data)]
        result["image_locations"] = ["Image 1"]

    else:
        result["images"] = [validate_raster(data)]
        result["image_locations"] = ["Image 1"]

    result["application_numbers"] = extract_application_numbers(
        result["text"]
    )

    result["comparison_text"] = remove_application_number_labels(
        result["text"]
    )

    return result



def load_pil_image(source: bytes | Image.Image) -> Image.Image:
    if isinstance(source, Image.Image):
        image = source.copy()
    else:
        image = Image.open(io.BytesIO(source))

    image.load()

    # Preserve transparency when identifying the foreground.
    rgba = image.convert("RGBA")
    background = Image.new("RGBA", rgba.size, (255, 255, 255, 255))
    background.alpha_composite(rgba)
    rgb = background.convert("RGB")

    # Crop away large white margins where possible.
    grayscale = rgb.convert("L")
    foreground_mask = grayscale.point(
        lambda pixel: 255 if pixel < 245 else 0
    )
    bounds = foreground_mask.getbbox()

    if bounds:
        left, top, right, bottom = bounds
        if right > left and bottom > top:
            rgb = rgb.crop(bounds)

    scale = 256 / max(rgb.width, rgb.height)
    rgb = rgb.resize(
        (
            max(1, round(rgb.width * scale)),
            max(1, round(rgb.height * scale)),
        ),
        Image.Resampling.LANCZOS,
    )
    return rgb


def prepare_comparison_image(source: bytes | Image.Image) -> Image.Image:
    logo = extract_logo_crop(source) if isinstance(source, bytes) else None
    return load_pil_image(logo or source)


def score_prepared_images(image_a: Image.Image, image_b: Image.Image) -> float:
    canvas_a = Image.new("RGB", (256, 256), "white")
    canvas_b = Image.new("RGB", (256, 256), "white")

    canvas_a.paste(
        image_a,
        ((256 - image_a.width) // 2, (256 - image_a.height) // 2),
    )
    canvas_b.paste(
        image_b,
        ((256 - image_b.width) // 2, (256 - image_b.height) // 2),
    )

    gray_a = canvas_a.convert("L").resize((64, 64))
    gray_b = canvas_b.convert("L").resize((64, 64))

    difference = ImageChops.difference(gray_a, gray_b)
    mean_difference = ImageStat.Stat(difference).mean[0]

    pixel_score = max(
        0.0,
        100.0 * (1.0 - mean_difference / 255.0),
    )

    hist_a = gray_a.histogram()
    hist_b = gray_b.histogram()

    hist_distance = sum(
        abs(left - right)
        for left, right in zip(hist_a, hist_b)
    )

    histogram_score = max(
        0.0,
        100.0 * (
            1.0 - hist_distance / (2.0 * 64 * 64)
        ),
    )

    mask_a = np.asarray(gray_a) < 237
    mask_b = np.asarray(gray_b) < 237
    union = np.count_nonzero(mask_a | mask_b)
    silhouette_score = (
        100.0 * np.count_nonzero(mask_a & mask_b) / union
        if union
        else pixel_score
    )

    return float(round(
        0.20 * pixel_score
        + 0.15 * histogram_score
        + 0.65 * silhouette_score,
        2,
    ))


def compare_images(
    source_a: bytes | Image.Image,
    source_b: bytes | Image.Image,
) -> float:
    return score_prepared_images(
        prepare_comparison_image(source_a),
        prepare_comparison_image(source_b),
    )


def compare_image_collections(
    images_a: list[bytes],
    images_b: list[bytes],
) -> dict[str, Any] | None:
    if not images_a or not images_b:
        return None

    prepared_a: list[tuple[int, Image.Image]] = []
    prepared_b: list[tuple[int, Image.Image]] = []

    for images, prepared in (
        (images_a, prepared_a),
        (images_b, prepared_b),
    ):
        for index, image in enumerate(images):
            try:
                prepared.append((index, prepare_comparison_image(image)))
            except Exception:
                continue

    best_match: dict[str, Any] | None = None
    similar_matches: list[dict[str, Any]] = []

    for index_a, image_a in prepared_a:
        for index_b, image_b in prepared_b:
            score = score_prepared_images(image_a, image_b)
            if score >= LOGO_MATCH_THRESHOLD:
                similar_matches.append({
                    "reference_image_index": index_a + 1,
                    "comparison_image_index": index_b + 1,
                    "similarity": score,
                })

            if best_match is None or score > best_match["similarity"]:
                best_match = {
                    "reference_image_index": index_a + 1,
                    "comparison_image_index": index_b + 1,
                    "similarity": score,
                }

    if best_match is not None:
        best_match["matches"] = sorted(
            similar_matches,
            key=lambda match: (
                -match["similarity"],
                match["reference_image_index"],
                match["comparison_image_index"],
            ),
        )
    return best_match

def detect_logo_region(source: bytes) -> dict[str, Any] | None:
    """
    Detect a likely foreground/logo bounding box in an extracted raster.
    Coordinates refer to the original decoded image, not a resized preview.
    This is a contour heuristic, not a trained logo detector.
    """
    array = np.frombuffer(source, dtype=np.uint8)
    image = cv2.imdecode(array, cv2.IMREAD_UNCHANGED)

    if image is None:
        return None

    if image.ndim == 2:
        gray = image
        height, width = gray.shape
    else:
        height, width = image.shape[:2]

        if image.shape[2] == 4:
            bgr = image[:, :, :3].astype(np.float32)
            alpha = image[:, :, 3:4].astype(np.float32) / 255.0
            image = (
                bgr * alpha + 255.0 * (1.0 - alpha)
            ).astype(np.uint8)
        elif image.shape[2] == 3:
            pass
        else:
            return None

        gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)

    if width < 16 or height < 16:
        return None

    image_area = width * height

    # Identify pixels that differ from a white background.
    foreground = cv2.threshold(
        gray, 237, 255, cv2.THRESH_BINARY_INV
    )[1]

    kernel_size = max(3, min(9, (min(width, height) // 100) | 1))
    kernel = np.ones((kernel_size, kernel_size), dtype=np.uint8)

    foreground = cv2.morphologyEx(
        foreground, cv2.MORPH_CLOSE, kernel
    )

    contours, _ = cv2.findContours(
        foreground,
        cv2.RETR_EXTERNAL,
        cv2.CHAIN_APPROX_SIMPLE,
    )

    candidates = []

    for contour in contours:
        x, y, box_width, box_height = cv2.boundingRect(contour)
        area = cv2.contourArea(contour)
        box_area = box_width * box_height

        if box_width < 12 or box_height < 12:
            continue

        if area < image_area * 0.003:
            continue

        # Reject full-canvas/background contours.
        if box_area > image_area * 0.92:
            continue

        candidates.append(
            (area, x, y, box_width, box_height)
        )

    if not candidates:
        return None

    area, x, y, box_width, box_height = max(
        candidates, key=lambda item: item[0]
    )

    return {
        "x": int(x),
        "y": int(y),
        "width": int(box_width),
        "height": int(box_height),
        "image_width": int(width),
        "image_height": int(height),
        "shape": "ellipse",
        "method": "opencv_contour_heuristic",
    }


def extract_logo_crop(source: bytes) -> Image.Image | None:
    region = detect_logo_region(source)
    if region is None:
        return None

    with Image.open(io.BytesIO(source)) as image:
        image.load()
        left = region["x"]
        top = region["y"]
        right = left + region["width"]
        bottom = top + region["height"]
        padding_x = max(2, int(region["width"] * 0.08))
        padding_y = max(2, int(region["height"] * 0.08))
        bounds = (
            max(0, left - padding_x),
            max(0, top - padding_y),
            min(image.width, right + padding_x),
            min(image.height, bottom + padding_y),
        )
        return image.crop(bounds)


def build_visual_evidence(
    images_a: list[bytes],
    images_b: list[bytes],
    image_match: dict[str, Any] | None,
) -> dict[str, Any]:
    regions: list[dict[str, Any]] = []

    if image_match is None:
        return {
            "regions": [],
            "max_regions": 4,
            "status": "NO_COMPARABLE_IMAGES",
        }

    matched_images = [
        (
            "reference",
            images_a,
            image_match["reference_image_index"],
        ),
        (
            "comparison",
            images_b,
            image_match["comparison_image_index"],
        ),
    ]

    for side, images, image_index in matched_images:
        index = image_index - 1

        if index < 0 or index >= len(images):
            continue

        try:
            region = detect_logo_region(images[index])
        except Exception:
            region = None

        if region is None:
            continue

        regions.append({
            "side": side,
            "image_index": image_index,
            **region,
        })

    return {
        "regions": regions[:4],
        "max_regions": 4,
        "status": (
            "REGIONS_DETECTED"
            if regions
            else "NO_REGIONS_DETECTED"
        ),
        "coordinate_space": "original_extracted_image_pixels",
        "method": "opencv_contour_heuristic",
    }

def classify_risk(score: float) -> str:
    if score >= 80:
        return "HIGH"
    if score >= 55:
        return "MEDIUM"
    return "LOW"


async def read_upload(file: UploadFile) -> bytes:
    filename = Path(file.filename or "").name

    if not filename or filename.startswith("."):
        raise HTTPException(
            status_code=400,
            detail="A valid filename is required.",
        )

    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=415,
            detail=(
                f"Unsupported file type: {extension or '(none)'}. "
                "Supported types: PDF, XLSX, XLSM, SVG, PNG, JPG, WEBP, TXT."
            ),
        )

    data = await file.read(MAX_FILE_BYTES + 1)

    if not data:
        raise HTTPException(
            status_code=400,
            detail=f"{filename} is empty.",
        )

    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"{filename} exceeds the 50 MB per-file limit.",
        )

    return data


@app.get("/")
def root() -> dict[str, str]:
    return {
        "name": APP_NAME,
        "version": VERSION,
        "status": "ready",
        "docs": "/docs",
    }


@app.get("/health")
def health() -> dict[str, Any]:
    return {
        "status": "ok",
        "service": "trademark-comparison-api",
        "version": VERSION,
        "auth_required": auth_is_required(),
        "community_access_managed": community_access_enabled(),
        "community_solution_code": os.getenv("COMMUNITY_SOLUTION_CODE", "trademark-conflict-detector"),
    }


@app.get("/api/v1/auth/config")
def auth_config() -> dict[str, Any]:
    return {
        "required": auth_is_required(),
        "verification": "email",
        "community_access_managed": community_access_enabled(),
        "community_solution_code": os.getenv("COMMUNITY_SOLUTION_CODE", "trademark-conflict-detector"),
    }


@app.post("/api/v1/auth/register")
def register(payload: AuthPayload) -> dict[str, str]:
    try:
        register_user(payload.email, payload.password)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except (OSError, smtplib.SMTPException) as exc:
        raise HTTPException(
            status_code=503,
            detail="Verification email could not be sent. Check SMTP configuration and try again.",
        ) from exc
    return {"status": "verification_sent"}


@app.post("/api/v1/auth/verify")
def verify(payload: EmailVerificationPayload) -> dict[str, str]:
    try:
        email = verify_email(payload.token)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"status": "verified", "email": email}


@app.post("/api/v1/auth/login")
def login(payload: AuthPayload) -> dict[str, str]:
    try:
        token = login_user(payload.email, payload.password)
    except ValueError as exc:
        raise HTTPException(status_code=401, detail=str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    return {"access_token": token, "token_type": "bearer"}


@app.post("/api/v1/compare")
async def compare_files(
    file_a: UploadFile = File(...),
    file_b: UploadFile = File(...),
    _user: str | None = Depends(require_user),
) -> dict[str, Any]:
    data_a = await read_upload(file_a)
    data_b = await read_upload(file_b)

    try:
        analyzed_a = analyze_file(file_a.filename or "file_a", data_a)
        analyzed_b = analyze_file(file_b.filename or "file_b", data_b)

    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    limitations = list(
        dict.fromkeys(
            analyzed_a["limitations"] + analyzed_b["limitations"]
        )
    )

    # Compare text after excluding labelled application-number fields.
    text_score = calculate_text_similarity(
        analyzed_a["comparison_text"],
        analyzed_b["comparison_text"],
    )

    image_match = compare_image_collections(
        analyzed_a["images"],
        analyzed_b["images"],
    )

    visual_evidence = build_visual_evidence(
        analyzed_a["images"],
        analyzed_b["images"],
        image_match,
    )
    logo_matches = image_match["matches"] if image_match is not None else []
    visual_match_summary = (
        {
            key: value
            for key, value in image_match.items()
            if key != "matches"
        }
        if image_match is not None
        else None
    )
    for match in logo_matches:
        reference_index = match["reference_image_index"] - 1
        comparison_index = match["comparison_image_index"] - 1
        match["reference_location"] = analyzed_a["image_locations"][
            reference_index
        ]
        match["comparison_location"] = analyzed_b["image_locations"][
            comparison_index
        ]

    if visual_evidence["status"] == "NO_REGIONS_DETECTED":
        limitations.append(
            "OpenCV could not localize a likely logo region in the "
            "selected images. No evidence outline was generated."
    )

    visual_score = (
        image_match["similarity"]
        if image_match is not None
        else None
    )

    if text_score is None:
        limitations.append(
            "Text similarity is unavailable because one or both files "
            "did not yield comparable text."
        )

    if visual_score is None:
        limitations.append(
            "No comparable raster images were extracted from both files. "
            "Visual similarity is unavailable, not zero."
        )

    # Never silently treat a missing visual score as 0% similarity.
    available_scores: list[tuple[float, float]] = []

    if text_score is not None:
        available_scores.append((text_score, 0.40 if visual_score is not None else 1.0))

    if visual_score is not None:
        available_scores.append((visual_score, 0.60 if text_score is not None else 1.0))

    if available_scores:
        total_weight = sum(weight for _, weight in available_scores)
        overall_score = round(
            sum(score * weight for score, weight in available_scores)
            / total_weight,
            2,
        )
        risk_level = classify_risk(overall_score)
        assessment_status = "PRELIMINARY"
    else:
        overall_score = None
        risk_level = "NOT_ASSESSED"
        assessment_status = "INSUFFICIENT_DATA"

    numbers_a = set(analyzed_a["application_numbers"])
    numbers_b = set(analyzed_b["application_numbers"])

    if numbers_a and numbers_b:
        number_status = (
            "MATCH" if numbers_a.intersection(numbers_b)
            else "DIFFERENT"
        )
    else:
        number_status = "NOT_DETECTED"

    limitations.append(
        "Application-number detection relies on labelled text patterns; "
        "formats vary by jurisdiction and OCR is not included."
    )
    limitations.append(
        "Logo visual scores use a basic pixel, histogram, and silhouette "
        "heuristic. They are not a trained logo-recognition model and may "
        "miss rotated, cropped, recolored, or semantically similar marks."
    )
    limitations.append(
        "This is a preliminary screening result, not a legal opinion, "
        "probability of infringement, or trademark-office decision."
    )

    return {
        "result": {
            "overall_similarity": overall_score,
            "text_similarity": text_score if text_score is not None else 0.0,
            "visual_similarity": visual_score if visual_score is not None else 0.0,
            "visual_status": (
                "AVAILABLE" if visual_score is not None else "NOT_AVAILABLE"
            ),
            "risk_level": risk_level,
            "assessment_status": assessment_status,
            "application_number": {
                "status": number_status,
                "reference_numbers": sorted(numbers_a),
                "comparison_numbers": sorted(numbers_b),
            },
            "visual_match": visual_match_summary,
            "logo_matches": logo_matches,
            "logo_match_threshold": LOGO_MATCH_THRESHOLD,
            "visual_evidence": visual_evidence,
            "limitations": list(dict.fromkeys(limitations)),
        },
        "files": [
            {
                "name": Path(file_a.filename or "file_a").name,
                "size_bytes": len(data_a),
                "extracted_image_count": len(analyzed_a["images"]),
            },
            {
                "name": Path(file_b.filename or "file_b").name,
                "size_bytes": len(data_b),
                "extracted_image_count": len(analyzed_b["images"]),
            },
        ],
    }


def _require_image_upload(filename: str) -> None:
    if Path(filename).suffix.lower() not in {
        ".png", ".jpg", ".jpeg", ".webp", ".svg",
    }:
        raise HTTPException(
            status_code=415,
            detail="Upload a PNG, JPG, WEBP, or SVG logo image.",
        )


@app.post("/api/v1/logo/crop")
async def crop_logo(
    file: UploadFile = File(...),
    _user: str | None = Depends(require_user),
) -> Response:
    from app.logo_search import extract_logo_crop

    filename = Path(file.filename or "").name
    _require_image_upload(filename)
    data = await read_upload(file)
    try:
        analyzed = analyze_file(filename, data)
        if not analyzed["images"]:
            raise ValueError("No image could be extracted from the upload.")
        original = analyzed["images"][0]
        region = detect_logo_region(original)
        cropped = extract_logo_crop(original)
        if cropped is None:
            raise ValueError("No logo region was detected in the uploaded image.")
        output = io.BytesIO()
        cropped.save(output, format="PNG", optimize=True)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    return Response(
        content=output.getvalue(),
        media_type="image/png",
        headers={
            "X-Logo-Region": json.dumps(region or {}),
            "Cache-Control": "no-store",
        },
    )


@app.post("/api/v1/trademarks")
async def add_trademark(
    image: UploadFile = File(...),
    jurisdiction: str = Form(...),
    mark_name: str = Form(...),
    trademark_number: str = Form(""),
    classes: str = Form(""),
    goods_services: str = Form(""),
    source_url: str = Form(""),
    x_ingest_key: str | None = Header(default=None),
    _user: str | None = Depends(require_user),
) -> dict[str, Any]:
    from app.logo_search import index_trademark

    configured_key = os.getenv("TRADEMARK_INGEST_KEY")
    if not configured_key:
        raise HTTPException(
            status_code=503,
            detail="Trademark ingestion is disabled until TRADEMARK_INGEST_KEY is configured.",
        )
    if not x_ingest_key or not hmac.compare_digest(x_ingest_key, configured_key):
        raise HTTPException(status_code=401, detail="Invalid ingestion credentials.")

    filename = Path(image.filename or "").name
    _require_image_upload(filename)
    data = await read_upload(image)
    try:
        analyzed = analyze_file(filename, data)
        if not analyzed["images"]:
            raise ValueError("No image could be extracted from the upload.")
        record = await run_in_threadpool(
            index_trademark,
            image_data=analyzed["images"][0],
            jurisdiction=jurisdiction,
            mark_name=mark_name,
            trademark_number=trademark_number,
            classes=classes,
            goods_services=goods_services,
            source_url=source_url,
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    return {"trademark": record}


@app.post("/api/v1/trademarks/search")
async def search_trademark_index(
    image: UploadFile = File(...),
    mark_name: str = Form(""),
    goods_services: str = Form(""),
    classes: str = Form(""),
    jurisdictions: list[str] = Form(default=[]),
    limit: int = Form(10),
    _user: str | None = Depends(require_user),
) -> dict[str, Any]:
    from app.logo_search import search_trademarks

    filename = Path(image.filename or "").name
    _require_image_upload(filename)
    data = await read_upload(image)
    try:
        analyzed = analyze_file(filename, data)
        if not analyzed["images"]:
            raise ValueError("No image could be extracted from the upload.")
        return await run_in_threadpool(
            search_trademarks,
            image_data=analyzed["images"][0],
            mark_name=mark_name,
            goods_services=goods_services,
            classes=classes,
            jurisdictions=jurisdictions,
            limit=limit,
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc


@app.get("/api/v1/trademarks/{trademark_id}/logo")
def trademark_logo(
    trademark_id: str,
    _user: str | None = Depends(require_user),
) -> Response:
    from app.logo_search import get_logo

    image = get_logo(trademark_id)
    if image is None:
        raise HTTPException(status_code=404, detail="Trademark logo not found.")
    return Response(image, media_type="image/png", headers={"Cache-Control": "public, max-age=3600"})


@app.get("/api/v1/trademarks/status")
def trademark_index_status() -> dict[str, Any]:
    from app.logo_search import get_index_count

    count = get_index_count()
    return {
        "indexed_trademarks": count,
        "jurisdictions": ["MYANMAR", "WIPO", "EUIPO", "TMVIEW", "USPTO"],
        "data_policy": "Only imported records are searchable; registry data is not fetched automatically.",
    }
